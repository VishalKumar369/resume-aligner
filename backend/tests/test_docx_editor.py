"""In-place .docx optimization keeps the user's formatting."""

from io import BytesIO

import docx
from docx.shared import RGBColor

from app.services.documents.docx_editor import (
    apply_rewrites_to_docx,
    optimize_docx_in_place,
)
from app.services.documents.keyword_highlight import compile_keyword_pattern


def _resume() -> bytes:
    document = docx.Document()
    name = document.add_paragraph()
    run = name.add_run("Vishal Kumar")
    run.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)  # navy

    bullet = document.add_paragraph("Built APIs in Python", style="List Bullet")
    bullet.runs[0].font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    document.add_paragraph("Led a team of five engineers", style="List Bullet")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _paragraphs(data: bytes):
    return [p for p in docx.Document(BytesIO(data)).paragraphs if p.text.strip()]


class TestApplyRewritesToDocx:
    def test_replaces_a_matched_bullet_keeping_its_formatting(self):
        edited, applied = apply_rewrites_to_docx(
            _resume(),
            [("Built APIs in Python", "Engineered REST APIs in Python, cutting p99 latency 30%")],
        )
        assert applied == 1

        paras = _paragraphs(edited)
        # The name paragraph is untouched: text, colour, and weight all kept.
        assert paras[0].text == "Vishal Kumar"
        assert paras[0].runs[0].bold is True
        assert paras[0].runs[0].font.color.rgb == RGBColor(0x1F, 0x4E, 0x79)
        # The bullet's wording changed but its list style and colour are kept.
        assert paras[1].text == "Engineered REST APIs in Python, cutting p99 latency 30%"
        assert paras[1].style.name == "List Bullet"
        assert paras[1].runs[0].font.color.rgb == RGBColor(0x33, 0x33, 0x33)
        # An unrelated bullet is left exactly as it was.
        assert paras[2].text == "Led a team of five engineers"

    def test_matches_ignoring_whitespace_and_case(self):
        _, applied = apply_rewrites_to_docx(
            _resume(), [("  built   apis in PYTHON ", "Shipped Python services")]
        )
        assert applied == 1

    def test_reports_zero_when_nothing_matches(self):
        # The caller uses this to fall back to the template renderer.
        _, applied = apply_rewrites_to_docx(
            _resume(), [("A bullet that isn't in the resume", "…")]
        )
        assert applied == 0

    def test_no_rewrites_returns_a_valid_unchanged_document(self):
        edited, applied = apply_rewrites_to_docx(_resume(), [])
        assert applied == 0
        assert [p.text for p in _paragraphs(edited)] == [
            "Vishal Kumar",
            "Built APIs in Python",
            "Led a team of five engineers",
        ]


class TestOptimizeDocxInPlace:
    def test_bolds_jd_keywords_in_bullets_only(self):
        pattern = compile_keyword_pattern(["Python", "APIs"])
        edited, applied, highlighted = optimize_docx_in_place(_resume(), [], pattern)

        assert applied == 0
        assert highlighted == 2  # "APIs" and "Python" in the bullet
        paras = _paragraphs(edited)
        # The name is left alone — not split into keyword runs.
        assert paras[0].text == "Vishal Kumar"
        assert len(paras[0].runs) == 1
        bolded = [r.text for r in paras[1].runs if r.bold]
        assert "Python" in bolded and "APIs" in bolded

    def test_highlights_the_rewritten_wording(self):
        # A rewrite runs first, so its new text is what gets highlighted.
        pattern = compile_keyword_pattern(["Kafka"])
        edited, applied, highlighted = optimize_docx_in_place(
            _resume(),
            [("Built APIs in Python", "Streamed events with Kafka at scale")],
            pattern,
        )
        assert applied == 1
        assert highlighted == 1
        bullet = _paragraphs(edited)[1]
        assert bullet.text == "Streamed events with Kafka at scale"
        assert [r.text for r in bullet.runs if r.bold] == ["Kafka"]
