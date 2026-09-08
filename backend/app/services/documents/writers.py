"""Export an optimized resume as plain text, .docx, and PDF.

All three render the same block list from `layout`, so they cannot drift. The
output is a polished, single-column professional resume: a centred name, contact
line with clickable links, section headings underlined by a rule, entry lines
with the role/company on the left and dates pushed to the right margin, and
JD-keyword bullets. Several `TEMPLATES` vary the colour, name size, and density
so the user can pick a look; the header rule and right-aligned dates use a light
table/tab, which modern ATS parse fine.

The .docx and the PDF share one page geometry and typography spec, and the .docx
uses *exact* line spacing, so a line is the same height in Word as in the PDF —
which is what lets the single-page condenser measure once and trust both.
"""

import re
from dataclasses import dataclass, replace
from io import BytesIO
from typing import Any, Dict, List, Optional, Pattern, Sequence, Tuple

from app.services.documents.keyword_highlight import segment_text
from app.services.documents.layout import BlockKind, build_blocks

# --------------------------------------------------------------- shared layout

# US Letter, in points (72pt = 1 inch). Both exporters use these exact values.
PAGE_WIDTH = 612.0
PAGE_HEIGHT = 792.0
MARGIN_TOP = 36.0      # 0.5"
MARGIN_BOTTOM = 36.0   # 0.5"
MARGIN_LEFT = 43.2     # 0.6"
MARGIN_RIGHT = 43.2    # 0.6"
USABLE_WIDTH = PAGE_WIDTH - MARGIN_LEFT - MARGIN_RIGHT

# Width of the right-hand (dates/location) column on an entry line.
RIGHT_COL_WIDTH = 150.0

# When *measuring* for the single-page condenser we reserve this much extra
# height, so a .docx that renders a hair taller than the PDF still lands on one
# page in Word. The real downloaded documents use the full page.
MEASURE_HEADROOM = 40.0  # ~0.55"

_LINK_COLOR = "#1155CC"
_RIGHT_COLOR = "#555555"


@dataclass(frozen=True)
class BlockStyle:
    size: float
    leading: float
    space_before: float = 0.0
    space_after: float = 0.0
    bold: bool = False
    italic: bool = False
    center: bool = False
    color: Optional[str] = None  # "#RRGGBB"


# The base typography. Templates are derived from this.
STYLE: Dict[BlockKind, BlockStyle] = {
    BlockKind.NAME: BlockStyle(18, 21, space_after=2, bold=True, center=True),
    BlockKind.CONTACT: BlockStyle(9, 12, space_after=4, center=True, color="#333333"),
    BlockKind.HEADING: BlockStyle(11, 14, space_before=9, space_after=2, bold=True),
    BlockKind.SUBHEADING: BlockStyle(10, 13, space_before=2, bold=True),
    BlockKind.META: BlockStyle(9, 12, space_after=1, italic=True, color="#444444"),
    BlockKind.PARAGRAPH: BlockStyle(9.5, 12.5, space_after=2),
    BlockKind.BULLET: BlockStyle(9.5, 12.5, space_after=1.5),
}


@dataclass(frozen=True)
class Template:
    """A selectable look: a typography map plus header-rule styling."""
    styles: Dict[BlockKind, BlockStyle]
    heading_rule: bool = True
    rule_color: str = "#333333"


def _with(overrides: Dict[BlockKind, BlockStyle]) -> Dict[BlockKind, BlockStyle]:
    merged = dict(STYLE)
    merged.update(overrides)
    return merged


_ACCENT = "#1F3A5F"  # a professional navy


TEMPLATES: Dict[str, Template] = {
    # Clean black-and-white with a rule under each heading. The safe default.
    "professional": Template(styles=dict(STYLE), heading_rule=True, rule_color="#333333"),
    # Navy accent on the name and headings, with a matching rule.
    "modern": Template(
        styles=_with({
            BlockKind.NAME: replace(STYLE[BlockKind.NAME], size=19, color=_ACCENT),
            BlockKind.HEADING: replace(STYLE[BlockKind.HEADING], color=_ACCENT),
        }),
        heading_rule=True,
        rule_color=_ACCENT,
    ),
    # Tighter type and spacing to pull a long resume onto one page.
    "compact": Template(
        styles={
            kind: replace(
                style,
                size=max(8.0, style.size - 0.5),
                leading=max(9.5, style.leading - 1.0),
                space_before=max(0.0, style.space_before - 2),
            )
            for kind, style in STYLE.items()
        },
        heading_rule=True,
        rule_color="#666666",
    ),
    # No rules — airy and understated. Headings carry weight on their own.
    "minimal": Template(
        styles=_with({
            BlockKind.NAME: replace(STYLE[BlockKind.NAME], size=17),
            BlockKind.HEADING: replace(STYLE[BlockKind.HEADING], space_before=11, space_after=3),
        }),
        heading_rule=False,
    ),
}

# "classic" is the pre-redesign id; keep it working as the professional default.
_ALIASES = {"classic": "professional", "original": "professional", "": "professional"}

# What the client may pick (besides "original", which keeps the uploaded file).
LAYOUT_IDS = ("professional", "modern", "compact", "minimal")


def _template(layout: Optional[str]) -> Template:
    key = (layout or "professional").lower()
    key = _ALIASES.get(key, key)
    return TEMPLATES.get(key, TEMPLATES["professional"])


# ------------------------------------------------------------------- plain text


def render_text(resume_data: Dict[str, Any], section_order: Optional[Sequence[str]] = None) -> str:
    """Plain-text rendering, used for keyword scoring of a generated resume."""
    lines: List[str] = []
    for block in build_blocks(resume_data, section_order):
        text = f"{block.text} {block.right}".strip() if block.right else block.text
        if block.kind is BlockKind.HEADING:
            lines.extend(["", text])
        elif block.kind is BlockKind.BULLET:
            lines.append(f"- {text}")
        else:
            lines.append(text)
    return "\n".join(lines).strip()


# ---------------------------------------------------------------------- linking

_URL_TOKEN = re.compile(r"^(?:https?://)?(?:www\.)?[A-Za-z0-9-]+\.[A-Za-z0-9./#?=&%+_-]+$")


def _as_url(token: str) -> Optional[str]:
    """The href for a contact token that is a link, else None (emails excluded)."""
    token = token.strip()
    if not token or "@" in token or "." not in token:
        return None
    if not _URL_TOKEN.match(token):
        return None
    return token if token.lower().startswith("http") else f"https://{token}"


def _contact_parts(text: str) -> List[str]:
    return [part.strip() for part in text.split("|") if part.strip()]


# ------------------------------------------------------------------------- docx


def render_docx(
    resume_data: Dict[str, Any],
    highlight_pattern: Optional[Pattern] = None,
    layout: Optional[str] = "professional",
    section_order: Optional[Sequence[str]] = None,
) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
    from docx.shared import Pt, RGBColor

    template = _template(layout)
    styles = template.styles

    def _rgb(hex_color: str) -> "RGBColor":
        return RGBColor(int(hex_color[1:3], 16), int(hex_color[3:5], 16), int(hex_color[5:7], 16))

    document = docx.Document()
    for section in document.sections:
        section.page_width = Pt(PAGE_WIDTH)
        section.page_height = Pt(PAGE_HEIGHT)
        section.top_margin = Pt(MARGIN_TOP)
        section.bottom_margin = Pt(MARGIN_BOTTOM)
        section.left_margin = Pt(MARGIN_LEFT)
        section.right_margin = Pt(MARGIN_RIGHT)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(styles[BlockKind.PARAGRAPH].size)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    def _apply_spacing(paragraph, spec: BlockStyle) -> None:
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(spec.space_before)
        fmt.space_after = Pt(spec.space_after)
        fmt.line_spacing = Pt(spec.leading)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY

    def _style_run(run, spec: BlockStyle, *, italic: Optional[bool] = None, color: Optional[str] = None) -> None:
        run.font.name = "Arial"
        run.font.size = Pt(spec.size)
        run.bold = spec.bold
        run.italic = spec.italic if italic is None else italic
        chosen = color if color is not None else spec.color
        if chosen:
            run.font.color.rgb = _rgb(chosen)

    def _highlight(paragraph) -> None:
        if highlight_pattern is not None:
            from app.services.documents.docx_editor import highlight_paragraph
            highlight_paragraph(paragraph, highlight_pattern)

    for block in build_blocks(resume_data, section_order):
        spec = styles.get(block.kind, styles[BlockKind.PARAGRAPH])

        if block.kind is BlockKind.CONTACT and "|" in block.text:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for i, part in enumerate(_contact_parts(block.text)):
                if i:
                    _style_run(paragraph.add_run("  |  "), spec)
                url = _as_url(part)
                if url:
                    _add_hyperlink(paragraph, url, part, _LINK_COLOR, spec.size)
                else:
                    _style_run(paragraph.add_run(part), spec)
            _apply_spacing(paragraph, spec)
            continue

        if block.kind is BlockKind.BULLET:
            paragraph = document.add_paragraph(block.text, style="List Bullet")
            for run in paragraph.runs:
                _style_run(run, spec)
            _apply_spacing(paragraph, spec)
            _highlight(paragraph)
            continue

        # A line with a right-aligned companion (dates/location): one run on the
        # left, a right tab stop, the companion after it.
        if block.right:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(USABLE_WIDTH), WD_TAB_ALIGNMENT.RIGHT)
            if block.text:
                _style_run(paragraph.add_run(block.text), spec)
            _style_run(paragraph.add_run("\t"), spec)
            _style_run(paragraph.add_run(block.right), spec, italic=True, color=_RIGHT_COLOR)
            _apply_spacing(paragraph, spec)
            continue

        paragraph = document.add_paragraph()
        _style_run(paragraph.add_run(block.text), spec)
        if spec.center:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _apply_spacing(paragraph, spec)

        if block.kind is BlockKind.HEADING and template.heading_rule:
            _set_bottom_border(paragraph, template.rule_color)
        elif block.kind is BlockKind.PARAGRAPH:
            # Summary and "Category: skills" lines carry JD keywords too.
            _highlight(paragraph)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_hyperlink(paragraph, url: str, text: str, color_hex: str, size_pt: float) -> None:
    from docx.oxml.ns import qn
    from docx.oxml.shared import OxmlElement

    part = paragraph.part
    r_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    props.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), color_hex.lstrip("#"))
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(int(size_pt * 2)))
    props.append(size)
    run.append(props)

    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _set_bottom_border(paragraph, color_hex: str) -> None:
    from docx.oxml.ns import qn
    from docx.oxml.shared import OxmlElement

    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color_hex.lstrip("#"))
    borders.append(bottom)
    p_pr.append(borders)


# -------------------------------------------------------------------------- pdf


def render_pdf(
    resume_data: Dict[str, Any],
    highlight_pattern: Optional[Pattern] = None,
    layout: Optional[str] = "professional",
    section_order: Optional[Sequence[str]] = None,
) -> bytes:
    payload, _pages = render_pdf_with_page_count(
        resume_data, highlight_pattern=highlight_pattern, layout=layout, section_order=section_order
    )
    return payload


def render_pdf_with_page_count(
    resume_data: Dict[str, Any],
    measure_headroom: float = 0.0,
    highlight_pattern: Optional[Pattern] = None,
    layout: Optional[str] = "professional",
    section_order: Optional[Sequence[str]] = None,
) -> Tuple[bytes, int]:
    """Render the PDF and report how many pages it occupies."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import (
        HRFlowable, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    template = _template(layout)
    tstyles = template.styles
    base = getSampleStyleSheet()["Normal"]

    def _para_style(kind: BlockKind, *, right: bool = False) -> ParagraphStyle:
        spec = tstyles.get(kind, tstyles[BlockKind.PARAGRAPH])
        italic = spec.italic or right
        font = "Helvetica-Bold" if (spec.bold and not right) else "Helvetica-Oblique" if italic else "Helvetica"
        style = ParagraphStyle(
            f"{kind.value}{'_r' if right else ''}", parent=base, fontName=font,
            fontSize=spec.size, leading=spec.leading,
            spaceBefore=0 if right else spec.space_before,
            spaceAfter=0 if right else spec.space_after,
        )
        if right:
            style.alignment = TA_RIGHT
            style.textColor = _RIGHT_COLOR
        else:
            if spec.center:
                style.alignment = TA_CENTER
            if spec.color:
                style.textColor = spec.color
        return style

    styles = {kind: _para_style(kind) for kind in tstyles}

    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer, pagesize=(PAGE_WIDTH, PAGE_HEIGHT),
        leftMargin=MARGIN_LEFT, rightMargin=MARGIN_RIGHT,
        topMargin=MARGIN_TOP, bottomMargin=MARGIN_BOTTOM + measure_headroom,
        title=str((resume_data.get("personal_info") or {}).get("name") or "Resume"),
    )

    story: List[Any] = []
    bullets: List[ListItem] = []

    def flush_bullets() -> None:
        if bullets:
            story.append(ListFlowable(
                list(bullets), bulletType="bullet", start="•",
                leftIndent=12, bulletFontSize=6, spaceAfter=2,
            ))
            bullets.clear()

    def entry_row(block) -> Table:
        spec = tstyles.get(block.kind, tstyles[BlockKind.PARAGRAPH])
        left = Paragraph(_escape(block.text), styles[block.kind]) if block.text else Spacer(0, 0)
        right = Paragraph(_escape(block.right), _para_style(block.kind, right=True))
        table = Table([[left, right]], colWidths=[USABLE_WIDTH - RIGHT_COL_WIDTH, RIGHT_COL_WIDTH])
        table.setStyle(TableStyle([
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), spec.space_before),
            ("BOTTOMPADDING", (0, 0), (-1, -1), spec.space_after),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        return table

    for block in build_blocks(resume_data, section_order):
        if block.kind is BlockKind.BULLET:
            markup = _highlight_markup(block.text, highlight_pattern)
            bullets.append(ListItem(Paragraph(markup, styles[BlockKind.BULLET]), leftIndent=12))
            continue

        flush_bullets()

        if block.kind is BlockKind.CONTACT and "|" in block.text:
            story.append(Paragraph(_contact_markup(block.text), styles[BlockKind.CONTACT]))
            continue

        if block.right:
            story.append(entry_row(block))
            continue

        if block.kind is BlockKind.PARAGRAPH:
            # Summary and "Category: skills" lines carry JD keywords too.
            story.append(Paragraph(_highlight_markup(block.text, highlight_pattern), styles[BlockKind.PARAGRAPH]))
            continue

        story.append(Paragraph(_escape(block.text), styles.get(block.kind, styles[BlockKind.PARAGRAPH])))
        if block.kind is BlockKind.HEADING and template.heading_rule:
            story.append(HRFlowable(
                width="100%", thickness=0.75, color=colors.HexColor(template.rule_color),
                spaceBefore=1, spaceAfter=3,
            ))

    flush_bullets()
    if not story:
        story.append(Spacer(1, 1))

    document.build(story)
    return buffer.getvalue(), max(1, int(getattr(document, "page", 1)))


def count_pdf_pages(
    resume_data: Dict[str, Any],
    headroom: float = MEASURE_HEADROOM,
    layout: Optional[str] = "professional",
    section_order: Optional[Sequence[str]] = None,
) -> int:
    """Pages the resume occupies (>= 1), with a safety headroom so the number is
    a floor the .docx honours too. Pass headroom=0 for the raw PDF page count."""
    return render_pdf_with_page_count(
        resume_data, measure_headroom=headroom, layout=layout, section_order=section_order
    )[1]


def _escape(text: str) -> str:
    """reportlab paragraphs accept inline markup, so raw text must be escaped."""
    return (
        str(text)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _contact_markup(text: str) -> str:
    """A centred contact line with the URL parts turned into clickable links."""
    pieces: List[str] = []
    for part in _contact_parts(text):
        url = _as_url(part)
        if url:
            pieces.append(f'<a href="{url}" color="{_LINK_COLOR}">{_escape(part)}</a>')
        else:
            pieces.append(_escape(part))
    return "  |  ".join(pieces)


def _highlight_markup(text: str, pattern: Optional[Pattern]) -> str:
    """Escaped reportlab markup with the JD's keyword matches wrapped in <b>."""
    if pattern is None:
        return _escape(text)
    return "".join(
        f"<b>{_escape(chunk)}</b>" if is_keyword else _escape(chunk)
        for chunk, is_keyword in segment_text(text, pattern)
    )
