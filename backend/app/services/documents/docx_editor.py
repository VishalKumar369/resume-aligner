"""Optimize a resume by editing the user's own .docx in place.

Rebuilding from structured data (``writers.py``) yields an ATS-clean but generic
document: it drops the colours, fonts, hyperlinks, and layout the user designed.
When the upload is a .docx we can instead edit that very file — replacing only
the bullet text the optimizer rewrote — so everything else the user styled
(header links, section colours, spacing) is preserved exactly.

A matching PDF is produced by converting the edited .docx with LibreOffice when
it is installed; callers fall back to the template renderer when it is not.
"""

import copy
import os
import re
import shutil
import subprocess
import tempfile
from io import BytesIO
from typing import List, Optional, Pattern, Tuple

from app.services.documents.keyword_highlight import segment_text
from app.services.parsing.line_utils import is_bullet


def _norm(text: str) -> str:
    """Whitespace- and case-insensitive key for matching a bullet to a paragraph."""
    return re.sub(r"\s+", " ", (text or "").strip()).lower()


def optimize_docx_in_place(
    original: bytes,
    rewrites: List[Tuple[str, str]],
    keyword_pattern: Optional[Pattern] = None,
) -> Tuple[bytes, int, int]:
    """Apply the bullet rewrites and highlight the JD keywords in one pass.

    Returns (edited_docx_bytes, replacements_applied, keywords_highlighted). The
    document is loaded once; rewrites run first (so a rewritten bullet's new
    wording is what gets highlighted).
    """
    import docx  # imported lazily; python-docx is a heavy import

    document = docx.Document(BytesIO(original))
    applied = _apply_rewrites(document, rewrites)
    highlighted = highlight_bullets_in_docx(document, keyword_pattern)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue(), applied, highlighted


def apply_rewrites_to_docx(
    original: bytes, rewrites: List[Tuple[str, str]]
) -> Tuple[bytes, int]:
    """Return (edited_docx_bytes, replacements_applied) — rewrites only."""
    import docx

    document = docx.Document(BytesIO(original))
    applied = _apply_rewrites(document, rewrites)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue(), applied


def _apply_rewrites(document, rewrites: List[Tuple[str, str]]) -> int:
    """Replace each paragraph whose text matches a rewrite's ``before``.

    Only matched paragraphs change; hyperlinks and styling elsewhere are left
    untouched. Returns how many replacements landed, so the caller can fall back
    to the template renderer when nothing matched.
    """
    wanted = {_norm(before): after for before, after in rewrites if before and after}
    if not wanted:
        return 0
    applied = 0
    for paragraph in _iter_paragraphs(document):
        replacement = wanted.get(_norm(paragraph.text))
        if replacement is not None and _replace_paragraph_text(paragraph, replacement):
            applied += 1
    return applied


def _iter_paragraphs(document):
    """Every paragraph in the body and in any tables (some resumes use tables)."""
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _replace_paragraph_text(paragraph, new_text: str) -> bool:
    """Set a paragraph's text to ``new_text`` while keeping its first run's
    formatting (font, size, weight, colour) and the paragraph's list style.

    The whole new string goes into the first run and the remaining runs are
    blanked, so the bullet keeps its look without duplicating mixed inline
    formatting we cannot faithfully re-segment.
    """
    runs = paragraph.runs
    if not runs:
        return False
    runs[0].text = new_text
    for run in runs[1:]:
        run.text = ""
    return True


def highlight_bullets_in_docx(document, pattern: Optional[Pattern]) -> int:
    """Bold every JD keyword occurrence inside the resume's bullet paragraphs.

    Scoped to bullets (list-styled or glyph-prefixed) so headings, the name, and
    the contact line are never touched. Returns the number of keywords bolded.
    """
    if pattern is None:
        return 0
    highlighted = 0
    for paragraph in _iter_paragraphs(document):
        if _is_body_bullet(paragraph):
            highlighted += _highlight_paragraph(paragraph, pattern)
    return highlighted


def _is_body_bullet(paragraph) -> bool:
    style_name = (paragraph.style.name or "") if paragraph.style is not None else ""
    return style_name.startswith("List") or is_bullet(paragraph.text)


def highlight_paragraph(paragraph, pattern: Optional[Pattern]) -> int:
    """Bold the keyword matches in one specific paragraph (used by the template
    renderer, which knows exactly which paragraphs are bullets/skills/summary)."""
    if pattern is None:
        return 0
    return _highlight_paragraph(paragraph, pattern)


def _highlight_paragraph(paragraph, pattern: Pattern) -> int:
    # Snapshot the runs: highlighting a run inserts new sibling runs after it,
    # which must not be re-processed.
    count = 0
    for run in list(paragraph.runs):
        count += _highlight_run(run, pattern)
    return count


def _highlight_run(run, pattern: Pattern) -> int:
    """Split one run so keyword matches become their own bold runs.

    The run's own formatting (font, size, colour, italic) is copied to every
    piece; only the weight of the keyword pieces changes. Returns how many
    keyword pieces were emphasised.
    """
    segments = segment_text(run.text, pattern)
    if not any(is_keyword for _, is_keyword in segments):
        return 0

    from docx.text.run import Run

    # Capture the original run's formatting before we mutate it, so every new
    # piece inherits the same look and only its bold flag varies.
    template = copy.deepcopy(run._element)
    base_bold = run.bold

    run.text = segments[0][0]
    run.bold = True if segments[0][1] else base_bold
    count = 1 if segments[0][1] else 0

    anchor = run._element
    for seg_text, is_keyword in segments[1:]:
        element = copy.deepcopy(template)
        anchor.addnext(element)
        anchor = element
        piece = Run(element, run._parent)
        piece.text = seg_text
        piece.bold = True if is_keyword else base_bold
        if is_keyword:
            count += 1
    return count


def docx_to_pdf(docx_bytes: bytes, timeout: float = 60.0) -> Optional[bytes]:
    """Convert .docx bytes to PDF bytes with LibreOffice; None if unavailable.

    Keeping the conversion here means the preserved-format .docx and its PDF are
    the same document, instead of the PDF silently falling back to the template.
    """
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None

    with tempfile.TemporaryDirectory() as tmp:
        source = os.path.join(tmp, "resume.docx")
        with open(source, "wb") as handle:
            handle.write(docx_bytes)
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, source],
                check=True,
                capture_output=True,
                timeout=timeout,
            )
        except (subprocess.SubprocessError, OSError):
            return None

        output = os.path.join(tmp, "resume.pdf")
        if not os.path.exists(output):
            return None
        with open(output, "rb") as handle:
            return handle.read()
