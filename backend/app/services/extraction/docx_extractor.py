import logging
from io import BytesIO
from typing import List, Optional, Tuple

logger = logging.getLogger(__name__)


def extract_docx(file_bytes: bytes) -> Tuple[str, Optional[str]]:
    """Extract text from a .docx, including tables.

    Many resume templates lay the whole document out in an invisible table, so
    skipping table cells would drop most of the content. Returns
    ``(text, error)``.
    """
    try:
        import docx
    except ImportError as exc:
        return "", f"python-docx unavailable: {exc}"

    try:
        document = docx.Document(BytesIO(file_bytes))
    except Exception as exc:  # noqa: BLE001 - malformed archives must not 500
        return "", str(exc)

    blocks: List[str] = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            # Word repeats the same cell object across merged spans.
            deduped = list(dict.fromkeys(cells))
            if deduped:
                blocks.append(" | ".join(deduped))

    blocks.extend(_extract_headers_and_footers(document))

    return "\n".join(blocks), None


def _extract_headers_and_footers(document) -> List[str]:
    """Contact details often live in the header, so they are worth recovering."""
    blocks: List[str] = []
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                if paragraph.text.strip():
                    blocks.append(paragraph.text.strip())
    return blocks
